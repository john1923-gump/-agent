from docx import Document
from ..models.schemas import Chapter
from ..utils.text_utils import gen_id, clean_text


def parse_docx(file_path: str, textbook_id: str, textbook_name: str) -> list[Chapter]:
    doc = Document(file_path)
    chapters: list[Chapter] = []
    current_title = ""
    current_content = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith('Heading') or _looks_like_heading(text):
            if current_content.strip():
                chapters.append(Chapter(
                    id=gen_id(),
                    textbook_id=textbook_id,
                    textbook_name=textbook_name,
                    title=current_title or f"第{len(chapters)+1}章",
                    content=clean_text(current_content),
                ))
            current_title = text
            current_content = ""
        else:
            current_content += text + "\n"

    if current_content.strip():
        chapters.append(Chapter(
            id=gen_id(),
            textbook_id=textbook_id,
            textbook_name=textbook_name,
            title=current_title or f"第{len(chapters)+1}章",
            content=clean_text(current_content),
        ))

    if not chapters:
        full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        chapters.append(Chapter(
            id=gen_id(),
            textbook_id=textbook_id,
            textbook_name=textbook_name,
            title="全文",
            content=clean_text(full_text),
        ))
    return chapters


def _looks_like_heading(text: str) -> bool:
    import re
    patterns = [
        r'^第[一二三四五六七八九十\d]+[章节篇]',
        r'^Chapter\s+\d+',
        r'^\d+[\.\s]+[A-Z\u4e00-\u9fff]',
    ]
    return bool(re.match('|'.join(patterns), text)) and len(text) < 60
