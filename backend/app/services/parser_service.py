import logging
import hashlib
from pathlib import Path
from ..models.schemas import Chapter, TextbookFormat, TextbookMeta
from ..utils.text_utils import gen_id
from ..parsers import pdf_parser, docx_parser, md_parser, txt_parser

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf": TextbookFormat.PDF,
    ".docx": TextbookFormat.DOCX,
    ".doc": TextbookFormat.DOCX,
    ".md": TextbookFormat.MD,
    ".markdown": TextbookFormat.MD,
    ".txt": TextbookFormat.TXT,
}

MAGIC_BYTES = {
    b"%PDF": TextbookFormat.PDF,
    b"PK\x03\x04": TextbookFormat.DOCX,
}

MIME_TYPE_MAP = {
    "application/pdf": TextbookFormat.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": TextbookFormat.DOCX,
    "application/msword": TextbookFormat.DOCX,
    "text/markdown": TextbookFormat.MD,
    "text/plain": TextbookFormat.TXT,
}


def _detect_by_magic_bytes(file_path: Path) -> TextbookFormat | None:
    try:
        with open(file_path, "rb") as f:
            header = f.read(16)
        for magic, fmt in MAGIC_BYTES.items():
            if header.startswith(magic):
                return fmt
    except Exception:
        pass
    return None


def _detect_by_mime_type(file_path: Path) -> TextbookFormat | None:
    try:
        import mimetypes
        mime_type, _ = mimetypes.guess_type(str(file_path))
        if mime_type:
            return MIME_TYPE_MAP.get(mime_type)
    except Exception:
        pass
    return None


def _validate_pdf(file_path: Path) -> str | None:
    try:
        import fitz
        doc = fitz.open(str(file_path))
        if doc.page_count == 0:
            return "PDF文件为空，无页面内容"
        text = ""
        for page in doc[:min(3, doc.page_count)]:
            text += page.get_text()
        doc.close()
        if len(text.strip()) < 50:
            return "PDF可能是扫描件，需要OCR处理才能提取文本"
        return None
    except Exception as e:
        return f"PDF文件损坏或格式无效: {str(e)}"


def _validate_docx(file_path: Path) -> str | None:
    try:
        from docx import Document
        doc = Document(str(file_path))
        if len(doc.paragraphs) == 0:
            return "DOCX文件无段落内容"
        text = "".join(p.text for p in doc.paragraphs[:10])
        if len(text.strip()) < 20:
            return "DOCX文件内容过少，可能为空文档"
        return None
    except Exception as e:
        return f"DOCX文件损坏或格式无效: {str(e)}"


def _validate_text(file_path: Path) -> str | None:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read(1000)
        if len(content.strip()) < 10:
            return "文本文件内容为空或过少"
        return None
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="gbk") as f:
                f.read(100)
            return None
        except Exception:
            return "文本文件编码不支持，请使用UTF-8或GBK编码"
    except Exception as e:
        return f"文本文件读取失败: {str(e)}"


def detect_format(filename: str, file_path: Path | None = None) -> tuple[TextbookFormat | None, str | None]:
    ext_fmt = SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower())

    if file_path and file_path.exists():
        magic_fmt = _detect_by_magic_bytes(file_path)
        if magic_fmt and ext_fmt and magic_fmt != ext_fmt:
            logger.warning(f"扩展名与内容不匹配: {filename} (扩展名={ext_fmt}, 内容={magic_fmt})")
            return magic_fmt, None

        mime_fmt = _detect_by_mime_type(file_path)
        if mime_fmt and ext_fmt and mime_fmt != ext_fmt:
            logger.warning(f"扩展名与MIME类型不匹配: {filename}")

    if ext_fmt:
        return ext_fmt, None

    if file_path and file_path.exists():
        magic_fmt = _detect_by_magic_bytes(file_path)
        if magic_fmt:
            return magic_fmt, None

    return None, f"不支持的文件格式: {Path(filename).suffix}，支持的格式: PDF, DOCX, MD, TXT"


def parse(file_path: Path, filename: str) -> tuple[list[Chapter], str | None]:
    fmt, detect_error = detect_format(filename, file_path)
    if detect_error:
        raise ValueError(detect_error)

    validation_error = None
    if fmt == TextbookFormat.PDF:
        validation_error = _validate_pdf(file_path)
    elif fmt == TextbookFormat.DOCX:
        validation_error = _validate_docx(file_path)
    else:
        validation_error = _validate_text(file_path)

    try:
        if fmt == TextbookFormat.PDF:
            chapters = pdf_parser.parse(file_path)
        elif fmt == TextbookFormat.DOCX:
            chapters = docx_parser.parse(file_path)
        elif fmt == TextbookFormat.MD:
            chapters = md_parser.parse(file_path)
        else:
            chapters = txt_parser.parse(file_path)
    except Exception as e:
        error_msg = _format_error_message(fmt, e, validation_error)
        raise ValueError(error_msg)

    if not chapters or all(len(ch.content.strip()) < 10 for ch in chapters):
        error_msg = _format_error_message(fmt, None, "解析结果为空或内容过少")
        raise ValueError(error_msg)

    return chapters, validation_error


def parse_textbook(file_path: str, filename: str) -> tuple[TextbookMeta, list[Chapter]]:
    path = Path(file_path)
    fmt, _ = detect_format(filename, path)

    try:
        chapters, validation_warning = parse(path, filename)
    except ValueError as e:
        logger.error(f"解析失败: {e}")
        raise

    content = "".join(ch.content for ch in chapters)
    import hashlib
    checksum = hashlib.md5(content.encode()).hexdigest()

    meta = TextbookMeta(
        id=gen_id(),
        filename=filename,
        format=fmt or TextbookFormat.TXT,
        chapter_count=len(chapters),
        checksum=checksum,
    )

    for ch in chapters:
        ch.textbook_id = meta.id

    if validation_warning:
        logger.warning(f"文件验证警告 ({filename}): {validation_warning}")

    return meta, chapters


def _format_error_message(fmt: TextbookFormat | None, exception: Exception | None, validation_error: str | None) -> str:
    parts = []
    if exception:
        parts.append(f"解析失败: {str(exception)}")
    if validation_error:
        parts.append(validation_error)
    if fmt == TextbookFormat.PDF:
        parts.append("提示: 如果是扫描版PDF，请先使用OCR工具转换")
    elif fmt == TextbookFormat.DOCX:
        parts.append("提示: 请确保DOCX文件未损坏且可正常打开")
    return " | ".join(parts) if parts else "文件解析失败"
